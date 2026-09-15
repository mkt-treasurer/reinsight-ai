"""Tests for the KB bordereau parser.

Two layers:
- The four real KB fixture files in ``tests/fixtures/kb_borderau`` are
  parsed end-to-end and asserted against known-good shape and totals.
- Synthetic in-memory workbooks exercise schema and edge-case handling
  (missing sheets, missing required columns, blank/decorative rows,
  ref_no parse failures) without polluting the fixtures directory.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.services.parsers.errors import KBSchemaError
from app.services.parsers.kb_bordereau import (
    KBBordereauResult,
    SHEET_KIND_PLA,
    SHEET_KIND_SOC,
    parse_kb_bordereau,
    parse_kb_sheet,
)


FIXTURES = Path(__file__).resolve().parents[2] / "fixtures" / "kb_borderau"

PLA_INS = FIXTURES / "PLA_FAC_INS_2026.02.xlsx"
PLA_DAEWOO = FIXTURES / "PLA_FAC_INS(DAEWOO)_2026.02.xlsx"
SOC_INS = FIXTURES / "SOC_FAC_INS_2026.02.xlsx"
SOC_DAEWOO = FIXTURES / "SOC_FAC_INS(DAEWOO)_2026.02.xlsx"


# ─── Real fixtures: top-level shape ────────────────────────────────────────


class TestRealFixtures:
    def test_pla_ins_basic_shape(self):
        r = parse_kb_bordereau(PLA_INS)
        assert r.sheet_kind == SHEET_KIND_PLA
        assert r.sheet_name == "OS List"
        assert "인스보험중개" in r.recipient
        assert len(r.rows) > 0
        assert r.skipped_sheets == []
        # All rows should successfully resolve a ref_no — none of the
        # fixture rows are sentinels or unparseable.
        assert all(row.ref_no is not None for row in r.rows)
        # PLA fixture is clean enough to produce zero warnings.
        assert r.warnings == []

    def test_pla_daewoo_recipient(self):
        r = parse_kb_bordereau(PLA_DAEWOO)
        assert r.recipient == "DAEWOO INS"
        assert r.sheet_kind == SHEET_KIND_PLA

    def test_soc_ins_basic_shape(self):
        r = parse_kb_bordereau(SOC_INS)
        assert r.sheet_kind == SHEET_KIND_SOC
        assert r.sheet_name == "SOC"
        assert "인스보험중개" in r.recipient
        assert len(r.rows) > 0

    def test_soc_daewoo_basic_shape(self):
        r = parse_kb_bordereau(SOC_DAEWOO)
        assert r.sheet_kind == SHEET_KIND_SOC
        assert r.recipient == "DAEWOO INS"
        assert len(r.rows) > 0


# ─── Real fixtures: per-row correctness ────────────────────────────────────


class TestPLARowExtraction:
    @pytest.fixture(scope="class")
    def result(self) -> KBBordereauResult:
        return parse_kb_bordereau(PLA_INS)

    def test_first_row_ref_no_normalised(self, result):
        first = result.rows[0]
        assert first.ref_no is not None
        # First medical-PI claim in the Feb 2026 PLA: ref 20241010024078
        # with 추산차수 0 (initial estimate).
        assert first.ref_no.normalized == "2024-1010024078 000"
        assert first.ref_no.year == 2024
        assert first.ref_no.month == 10
        assert first.ref_no.day == 10

    def test_first_row_business_fields(self, result):
        first = result.rows[0]
        assert first.insured == "의료법인목포구암의료재단목포중앙병원"
        assert first.dol == datetime(2024, 7, 4).date()
        assert first.doc == datetime(2024, 10, 10).date()
        assert first.currency_claim == "WON"
        assert first.your_share_pct == Decimal("30")

    def test_amounts_use_decimal(self, result):
        first = result.rows[0]
        assert first.amounts["os_loss_indemnity"] == Decimal("10000000")
        assert first.amounts["os_expense"] == Decimal("1200000")
        assert first.amounts["ceded_os_loss_indemnity"] == Decimal("3000000")

    def test_lob_forward_filled(self, result):
        # KB merges 상품명 (LoB) across consecutive rows; downstream code
        # needs the value populated on every row.
        for row in result.rows:
            assert row.lob is not None and row.lob.strip() != ""


class TestSiblingCasualties:
    """KB groups personal-injury (`인`) and property-damage (`물`) lines
    of the same incident under the same KB Ref + revision but distinct
    ``casualty`` values. These are *sibling* rows, not duplicates —
    both must survive parsing with distinct natural keys.

    Anchor: PLA INS Feb 2026 fixture, ``2023-0707020197`` rev 0
    (Paris hotel robbery, single insured ``(주) 온라인투어``):
        - row R78: casualty=``인001 미상``, OS_Loss=1,000,000
        - row R79: casualty=``물001 미상``, OS_Loss=10,000,000
    Same ref / same revision / same item & cession serials, but the
    casualty discriminator separates the personal vs. property arms of
    the same incident.
    """

    def test_pla_ins_preserves_personal_and_property_damage_as_separate_rows(self):
        r = parse_kb_bordereau(PLA_INS)
        siblings = [
            row for row in r.rows
            if row.ref_no is not None
            and row.ref_no.key == "2023-0707020197-000"
        ]
        assert len(siblings) == 2, (
            "expected exactly two sibling rows under "
            "(2023-0707020197, rev=0) in PLA INS Feb 2026"
        )

        casualties = sorted(row.casualty for row in siblings)
        assert casualties == ["물001 미상", "인001 미상"], (
            f"expected 인/물 sibling pair, got {casualties}"
        )

        # Both rows share item & cession serial; only casualty separates.
        assert {row.item_serial for row in siblings} == {1}
        assert {row.cession_serial for row in siblings} == {1}

        # Natural keys must differ — that's what makes them not duplicates.
        keys = {row.natural_key() for row in siblings}
        assert len(keys) == 2

    def test_natural_key_includes_casualty(self):
        r = parse_kb_bordereau(PLA_INS)
        siblings = [
            row for row in r.rows
            if row.ref_no is not None
            and row.ref_no.key == "2023-0707020197-000"
        ]
        # Strip casualty from each natural key; the rest of the tuple
        # should be identical between the two siblings.
        without_casualty = {
            (k[0], k[1], k[3], k[4], k[5])
            for k in (row.natural_key() for row in siblings)
        }
        assert len(without_casualty) == 1, (
            "natural keys differ on something other than casualty — "
            "the casualty field is supposed to be the sole discriminator"
        )

    def test_natural_key_uniqueness_across_full_fixtures(self):
        """No two rows across the four real fixtures should share a
        natural key. If this ever fires, the fixture has true duplicates
        and Phase-2 dedup policy needs to make an explicit decision."""
        from collections import Counter

        seen: Counter = Counter()
        for path in (PLA_INS, PLA_DAEWOO, SOC_INS, SOC_DAEWOO):
            r = parse_kb_bordereau(path)
            for row in r.rows:
                seen[row.natural_key()] += 1

        collisions = {k: n for k, n in seen.items() if n > 1}
        assert not collisions, (
            "natural-key collisions across fixtures — the parser is "
            "flattening rows that should be distinct, or the source "
            "data contains true duplicates that need a policy. "
            f"collisions: {collisions}"
        )


class TestSOCNegativesAndTotals:
    """Ensure negative reversal rows are preserved and the per-currency
    Total row matches the sum across data rows."""

    @pytest.fixture(scope="class")
    def soc_ins(self) -> KBBordereauResult:
        return parse_kb_bordereau(SOC_INS)

    @pytest.fixture(scope="class")
    def soc_daewoo(self) -> KBBordereauResult:
        return parse_kb_bordereau(SOC_DAEWOO)

    def test_soc_ins_won_total_matches_row_sum(self, soc_ins):
        won_total = soc_ins.totals_by_currency["WON"]
        row_sum_indemnity = sum(
            (row.amounts.get("paid_loss_indemnity", Decimal(0)) for row in soc_ins.rows),
            Decimal(0),
        )
        row_sum_others = sum(
            (row.amounts.get("paid_loss_others", Decimal(0)) for row in soc_ins.rows),
            Decimal(0),
        )
        row_sum_expense = sum(
            (row.amounts.get("paid_expense", Decimal(0)) for row in soc_ins.rows),
            Decimal(0),
        )
        assert won_total["paid_loss_indemnity"] == row_sum_indemnity
        assert won_total["paid_loss_others"] == row_sum_others
        assert won_total["paid_expense"] == row_sum_expense

    def test_soc_daewoo_total_preserves_negatives(self, soc_daewoo):
        # The Feb 2026 DAEWOO SOC has reversal rows pushing two amount
        # columns negative; document totals must come through with the
        # right sign.
        won_total = soc_daewoo.totals_by_currency["WON"]
        assert won_total["paid_loss_others"] == Decimal("-54565657")
        assert won_total["paid_expense"] == Decimal("-4653416")
        assert won_total["ceded_paid_loss_others"] == Decimal("-5456566")

    def test_soc_ins_negative_reversal_row_preserved(self, soc_ins):
        # The Bicos Smile Tour reversal row (2024-1105009358 002) should
        # keep its -5,000,000 paid_loss_indemnity.
        negatives = [
            row for row in soc_ins.rows
            if row.amounts.get("paid_loss_indemnity", Decimal(0)) < 0
        ]
        assert len(negatives) >= 1
        assert any(
            row.ref_no is not None
            and "2024-1105009358" in row.ref_no.normalized
            and row.amounts["paid_loss_indemnity"] == Decimal("-5000000")
            for row in negatives
        )


# ─── Synthetic workbooks: error and edge cases ─────────────────────────────


def _save_xlsx(wb: Workbook, tmp_path: Path, name: str = "test.xlsx") -> Path:
    """Persist an in-memory workbook to a temp file the parser can load."""
    path = tmp_path / name
    wb.save(path)
    return path


def _build_pla_header() -> list[str]:
    """Mirror the canonical KB PLA header row 12 layout."""
    return [
        None,  # col 1 placeholder for cedant tag column
        "상품명\n(LoB)",
        "증권번호 \n(Policy No)",
        "계약자명\n(Insured)",
        "보험개시일\n(Inception)",
        "보험종료일\n(Expiry)",
        "접수번호\n(KB Ref)",
        "사고일자\n(DOL)",
        "배상청구일자\n(DOC)",
        "사고내용\n(Loss Description)",
        "재해명\n(Casualty)",
        "추산일자\n(OS Date)",
        "추산차수\n(OS Serial)",
        "목적물번호\n(Item Serial)",
        "출재단위번호\n(Cession Serial)",
        "담보코드\n(Coverage Code)",
        "화폐\n(Currency_Claim)",
        "화폐\n(Currency_Expense)",
        "원보험금\n(OS Loss_Indemnity)",
        "원보험금_기타\n(OS Loss_Others)",
        "원 손해조사비\n(OS Expense)",
        "출재율(%)\nYour Share(%)",
        "재보험금\n(Ceded OS Loss_Indemnity)",
        "재보험금_기타\n(Ceded OS Loss_Others)",
        "출재 손해조사비\n(Ceded OS Expense)",
    ]


def _build_minimal_pla(sheet_title: str = "OS List") -> Workbook:
    """A KB-PLA-shaped workbook with letterhead + header + one row."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    # Letterhead rows 1-9
    for _ in range(9):
        ws.append([None])
    # Recipient row 10
    ws.append([None, "To :", "DAEWOO INS"])
    # Spacer row 11
    ws.append([None])
    # Header row 12
    ws.append(_build_pla_header())
    # One data row 13
    ws.append([
        "DAEWOO INS",
        "Commercial General Liability",
        "P-001",
        "테스트피보험자",
        datetime(2024, 1, 1),
        datetime(2025, 1, 1),
        "20240101000123",
        datetime(2024, 6, 1),
        datetime(2024, 6, 5),
        "loss desc",
        "casualty",
        datetime(2024, 6, 10),
        1,  # OS Serial
        1,
        1,
        "GLC011",
        "WON",
        "WON",
        100000,
        0,
        5000,
        10,
        10000,
        0,
        500,
    ])
    return wb


class TestSchemaErrors:
    def test_no_whitelisted_sheet_raises(self, tmp_path):
        wb = Workbook()
        wb.active.title = "RandomSheet"
        path = _save_xlsx(wb, tmp_path)
        with pytest.raises(KBSchemaError, match="no whitelisted KB sheet"):
            parse_kb_bordereau(path)

    def test_header_row_missing_raises(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "OS List"
        # Plenty of letterhead but never the 접수번호 marker.
        for i in range(15):
            ws.append([f"line {i}", None, None])
        path = _save_xlsx(wb, tmp_path)
        with pytest.raises(KBSchemaError, match="접수번호"):
            parse_kb_bordereau(path)

    def test_required_column_missing_raises(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "OS List"
        for _ in range(11):
            ws.append([None])
        # Header with 접수번호 but missing required share/insured/dol.
        ws.append([None, "접수번호\n(KB Ref)", "추산차수\n(OS Serial)"])
        ws.append([None, "20240101000123", 1])
        path = _save_xlsx(wb, tmp_path)
        with pytest.raises(KBSchemaError, match="required columns missing"):
            parse_kb_bordereau(path)


class TestRowHandling:
    def test_minimal_pla_parses(self, tmp_path):
        path = _save_xlsx(_build_minimal_pla(), tmp_path)
        r = parse_kb_bordereau(path)
        assert r.sheet_kind == SHEET_KIND_PLA
        assert r.recipient == "DAEWOO INS"
        assert len(r.rows) == 1
        first = r.rows[0]
        assert first.ref_no is not None
        assert first.ref_no.normalized == "2024-0101000123 001"
        assert first.amounts["os_loss_indemnity"] == Decimal("100000")

    def test_blank_row_silently_skipped(self, tmp_path):
        wb = _build_minimal_pla()
        ws = wb.active
        # Insert a fully blank row between header and the data row.
        ws.insert_rows(13)
        path = _save_xlsx(wb, tmp_path)
        r = parse_kb_bordereau(path)
        assert len(r.rows) == 1

    def test_decorative_row_silently_skipped(self, tmp_path):
        wb = _build_minimal_pla()
        ws = wb.active
        # A row with only a casualty label (no ref_no, insured, policy_no)
        # should be discarded as decorative.
        ws.append([None, None, None, None, None, None, None, None, None, None, "decorative"])
        path = _save_xlsx(wb, tmp_path)
        r = parse_kb_bordereau(path)
        assert len(r.rows) == 1

    def test_unparseable_ref_no_kept_with_warning(self, tmp_path):
        wb = _build_minimal_pla()
        ws = wb.active
        # Add a second data row with a clearly malformed ref_no — keep
        # the row (so totals reconcile) but record the warning.
        ws.append([
            None,
            "Commercial General Liability",
            "P-002",
            "다른피보험자",
            datetime(2024, 1, 1),
            datetime(2025, 1, 1),
            "NOT-A-REF",
            datetime(2024, 6, 1),
            None,
            "loss desc",
            "c",
            datetime(2024, 6, 10),
            1, 1, 1,
            "GLC011",
            "WON", "WON",
            50000, 0, 0, 5, 0, 0, 0,
        ])
        path = _save_xlsx(wb, tmp_path)
        r = parse_kb_bordereau(path)
        assert len(r.rows) == 2
        bad_row = r.rows[1]
        assert bad_row.ref_no is None
        assert any("ref_no parse failed" in note for note in bad_row.notes)
        assert any("NOT-A-REF" in w or "ref_no parse failed" in w for w in r.warnings)


class TestSheetSelection:
    def test_extra_sheets_recorded_as_skipped(self, tmp_path):
        wb = _build_minimal_pla()
        wb.create_sheet("Cover Page")
        wb.create_sheet("Notes")
        path = _save_xlsx(wb, tmp_path)
        r = parse_kb_bordereau(path)
        assert set(r.skipped_sheets) == {"Cover Page", "Notes"}

    def test_first_whitelisted_sheet_wins_when_both_present(self, tmp_path):
        """If a workbook ever ships both OS List and SOC, parser picks
        whichever appears first and records the other as skipped."""
        wb = _build_minimal_pla(sheet_title="OS List")
        soc = wb.create_sheet("SOC")
        soc.append([None])  # noise so the sheet exists
        path = _save_xlsx(wb, tmp_path)
        r = parse_kb_bordereau(path)
        assert r.sheet_kind == SHEET_KIND_PLA
        assert "SOC" in r.skipped_sheets

    def test_sheet_order_pla_first_means_pla_wins(self, tmp_path):
        wb = Workbook()
        wb.active.title = "Cover Page"
        # Add an OS List then a SOC. PLA should be picked because it
        # appears first in workbook order.
        pla_ws = wb.create_sheet("OS List")
        # Recreate minimal PLA structure on the new sheet.
        for _ in range(9):
            pla_ws.append([None])
        pla_ws.append([None, "To :", "DAEWOO INS"])
        pla_ws.append([None])
        pla_ws.append(_build_pla_header())
        pla_ws.append([
            "DAEWOO INS", "LoB", "P", "I",
            datetime(2024, 1, 1), datetime(2025, 1, 1),
            "20240101000123",
            datetime(2024, 6, 1), datetime(2024, 6, 5),
            "d", "c", datetime(2024, 6, 10),
            1, 1, 1, "GLC011", "WON", "WON",
            100000, 0, 5000, 10, 10000, 0, 500,
        ])
        wb.create_sheet("SOC").append([None])
        path = _save_xlsx(wb, tmp_path)
        r = parse_kb_bordereau(path)
        assert r.sheet_kind == SHEET_KIND_PLA
        assert "Cover Page" in r.skipped_sheets
        assert "SOC" in r.skipped_sheets


class TestParseKBSheetDirect:
    def test_passes_skipped_sheets_through(self):
        # parse_kb_sheet itself doesn't know about other sheets in the
        # workbook; the orchestrator is responsible for filling
        # skipped_sheets. Verify the direct entry point returns [].
        wb = load_workbook(PLA_INS, data_only=True)
        ws = wb["OS List"]
        r = parse_kb_sheet(ws, file_name=PLA_INS.name)
        assert r.skipped_sheets == []
        assert r.sheet_kind == SHEET_KIND_PLA

    def test_rejects_non_whitelisted_sheet(self):
        wb = Workbook()
        wb.active.title = "RandomSheet"
        with pytest.raises(KBSchemaError, match="not on the KB whitelist"):
            parse_kb_sheet(wb.active, file_name="random.xlsx")
