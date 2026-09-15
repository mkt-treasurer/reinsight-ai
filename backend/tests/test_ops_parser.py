"""Unit tests for `services.ops_parser`.

Pure logic — no DB, no FastAPI. Fixtures are built in-memory with openpyxl that
mirror the real workbook's two sheet families (Actual* contract ledger + bare-
year 선수금 settlement ledger), so no PII-bearing sample file is needed.
"""

from __future__ import annotations

from datetime import date

import openpyxl
import pytest

from app.services.ops_parser import (
    is_open_settlement,
    open_settlement,
    parse_workbook,
    summarise_settlement,
)

# ── fixtures ─────────────────────────────────────────────────────────────────

# Subset of the Actual* header in real column order.
_ACTUAL_HEADER = [
    "Cont. Month", "No.", "Cover Note No.", "Assured / Insured", "Line",
    "New/Renew", "Cedant", "From", "To", "PPW", "Currency",
    "Gross Prem.(100%)", "Reinsurer", "Share", "Rec. Date", "Paid Date",
    "Account Mgr.",
]


def _actual_row(**kw):
    base = {
        "Cont. Month": None, "No.": None, "Cover Note No.": None,
        "Assured / Insured": None, "Line": None, "New/Renew": None,
        "Cedant": None, "From": None, "To": None, "PPW": None, "Currency": None,
        "Gross Prem.(100%)": None, "Reinsurer": None, "Share": None,
        "Rec. Date": None, "Paid Date": None, "Account Mgr.": None,
    }
    base.update(kw)
    return [base[h] for h in _ACTUAL_HEADER]


def _build_workbook():
    wb = openpyxl.Workbook()
    # Actual2026 contract ledger
    ws = wb.active
    ws.title = "Actual2026"
    ws.append(["AI 계약관리"])  # a stray summary row above the header
    ws.append(_ACTUAL_HEADER)
    # one contract with two reinsurer shares (continuation row blanks identity)
    ws.append(_actual_row(**{
        "Cover Note No.": "CN-001", "Assured / Insured": "Acme Corp",
        "Cedant": "Samsung F&M", "Reinsurer": "AON", "Share": 60,
        "From": date(2026, 1, 1), "To": date(2026, 12, 31),
        "PPW": date(2026, 2, 28), "Gross Prem.(100%)": 1000000,
        "Account Mgr.": "Brian",
    }))
    ws.append(_actual_row(Reinsurer="Lockton", Share=40))  # continuation share

    # bare-year settlement sheet "2025"
    s = wb.create_sheet("2025")
    s.append([None, 154745, 105.62, 1465.1])  # totals row (no counterparty)
    s.append(["선수금입금내역", None, None, None, None, "선수금송금내역"])
    s.append(["Date", "입금액", None, None, "Reinsurer", "Date", "Amount",
              None, None, "수수료", None, "선수금잔액", "환차손/익", "Net to",
              "Assured", "Remarks", "PRE\n/CLM", "Account\nMgr."])
    s.append([None, "KRW", "USD", "ROE", None, None, "KRW", "USD", "ROE",
              "KRW", "USD"])
    # open: received, not remitted, nonzero balance, FX gain marker
    s.append([date(2025, 1, 7), 419632, 354.56, 1183.5, "AON", None, None,
              None, None, None, None, 419632, "환차익", "INS", "Acme Corp",
              "1Q2025", "PRE", "류제은"])
    # settled: received and remitted, zero balance, FX loss marker
    s.append([date(2025, 1, 7), 100000, 80.0, 1250.0, "Lockton",
              date(2025, 1, 10), 100000, 80.0, 1250.0, 0, 0, 0, "환차손",
              "INS", "Beta Ltd", "2Q2025", "PRE", "박영선"])
    return wb


@pytest.fixture()
def workbook_path(tmp_path):
    path = tmp_path / "AI계약관리.xlsx"
    _build_workbook().save(path)
    return str(path)


# ── contract ledger ──────────────────────────────────────────────────────────

def test_parses_actual_shares_with_forward_filled_identity(workbook_path):
    out = parse_workbook(workbook_path)

    assert out["sheets"] == {"Actual2026": 2}
    rows = out["rows"]
    assert len(rows) == 2
    # continuation share inherits identity but keeps its own reinsurer/share
    assert rows[1]["cover_note"] == "CN-001"
    assert rows[1]["assured"] == "Acme Corp"
    assert rows[1]["reinsurer"] == "Lockton"
    assert rows[1]["share"] == 40.0
    assert rows[0]["ppw"] == "2026-02-28"
    assert rows[0]["book_year"] == 2026


# ── settlement ledger ─────────────────────────────────────────────────────────

def test_parses_settlement_rows_and_skips_totals(workbook_path):
    out = parse_workbook(workbook_path)

    assert out["settlement_sheets"] == {"2025": 2}  # totals/header rows excluded
    s = out["settlement"]
    assert len(s) == 2
    first = s[0]
    assert first["recv_date"] == "2025-01-07"
    assert first["recv_krw"] == 419632.0
    assert first["recv_usd"] == 354.56
    assert first["reinsurer"] == "AON"
    assert first["balance"] == 419632.0
    assert first["assured"] == "Acme Corp"
    assert first["account_mgr"] == "류제은"


def test_fx_marker_stays_text_not_number(workbook_path):
    s = parse_workbook(workbook_path)["settlement"]
    assert s[0]["fx_gl"] == "환차익"
    assert s[1]["fx_gl"] == "환차손"


def test_is_open_settlement_distinguishes_pending_from_settled(workbook_path):
    s = parse_workbook(workbook_path)["settlement"]
    received_not_remitted, settled = s[0], s[1]
    assert is_open_settlement(received_not_remitted) is True
    assert is_open_settlement(settled) is False


def test_open_settlement_projects_to_ship_fields_only(workbook_path):
    from app.services.ops_parser import SETTLE_SHIP_FIELDS

    s = parse_workbook(workbook_path)["settlement"]
    open_rows = open_settlement(s)
    assert len(open_rows) == 1
    # Payload is trimmed to exactly the UI columns — no comm/roe/fx/remarks etc.
    assert set(open_rows[0].keys()) == set(SETTLE_SHIP_FIELDS)
    assert open_rows[0]["recv_krw"] == 419632.0
    assert open_rows[0]["reinsurer"] == "AON"


def test_summarise_settlement_rolls_up_currency_split_and_fx_counts(workbook_path):
    s = parse_workbook(workbook_path)["settlement"]
    summary = summarise_settlement(s)

    assert summary["row_count"] == 2
    assert summary["open_count"] == 1
    assert summary["recv_krw"] == 519632.0  # 419632 + 100000, KRW only
    assert summary["recv_usd"] == 434.56     # 354.56 + 80.0, never mixed with KRW
    assert summary["balance_krw"] == 419632.0
    assert summary["fx_gain"] == 1
    assert summary["fx_loss"] == 1
    assert "2025" in summary["by_year"]


def test_years_helper_returns_distinct_years_descending():
    from app.routers.ops import _years

    rows = [{"book_year": 2024}, {"book_year": 2026}, {"book_year": 2024},
            {"book_year": None}, {"book_year": 2025}]
    assert _years(rows) == [2026, 2025, 2024]
    assert _years([]) == []


def test_non_settlement_sheet_names_are_ignored():
    wb = openpyxl.Workbook()
    wb.active.title = "Notes"
    wb.active.append(["nothing", "useful"])
    import tempfile
    import os
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(path)
        out = parse_workbook(path)
        assert out["rows"] == []
        assert out["settlement"] == []
    finally:
        os.unlink(path)
