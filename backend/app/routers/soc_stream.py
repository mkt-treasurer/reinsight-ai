"""SSE streaming SOC generator — shows CoT in real-time."""

import asyncio
import json
import tempfile
import logging
import re
from pathlib import Path
from typing import AsyncGenerator

from fastapi import APIRouter, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
import google.generativeai as genai

from app.database import async_session
from app.config import settings
from app.models.contract import Contract
from app.models.claim import Claim
from app.models.policy import Policy
from app.models.slip_case import SlipCase
from app.services.reinsurer_names import resolve_full_name
from app.services.soc_prompts import (
    EXTRACT_CLAIM_BORDEREAU,
    REINSURER_REASONING_SHORT,
    SLIP_TRANSLATE_TO_ENGLISH,
)
from app.data.code_table import is_foreign_reinsurer, canonicalize_company_name
from app.services.soc_utils import (
    parse_json as _parse_json_raw,
    fmt_amount as _fmt_amount,
    to_num as _to_num_util,
    is_broker_name,
    normalize_cedant,
    strip_corp_suffix,
    claim_to_dict as _claim_to_dict,
    contract_to_dict as _contract_to_dict,
    fmt_claim as _fmt_claim,
    fmt_contract as _fmt_contract,
    resolve_insured_name,
    pick_english_insured_name,
)
from app.services.soc_stream_kb_integration import (
    finalize_kb_integration,
    init_kb_integration,
    run_kb_cross_checks,
    try_parse_kb_xlsx,
)
from app.services.parsers.kb import SEVERITY_FAIL, SEVERITY_WARNING
from app.services.parsers.errors import RefNoParseError
from app.services.parsers.ref_no_parser import parse_kb_ref_no
import uuid as _uuid


_MONTHS_EN_SHORT = (
    "Jan", "Feb", "Mar", "Apr", "May", "Jun",
    "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
)


def _to_english_date_range(s: str | None) -> str:
    """Normalize a policy-period string to the KB ``Mmm DD, YYYY`` form.

    Accepts ISO-ish dates (``2014-09-01`` / ``2014/09/01`` / ``2014.09.01``)
    anywhere in the string and rewrites them in place; range separators
    (``~`` / ``to`` / ``-``) pass through untouched. Strings that are
    already English (KB native ``Sep 01, 2014 ~ Aug 31, 2015``) or non-
    parseable (``Refer to CN12345``, ``TBD``) come back unchanged.

    Idempotent: re-running on already-English output is a no-op.
    """
    if not s:
        return ""
    out = str(s)

    def _sub(m: re.Match[str]) -> str:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if not 1 <= mo <= 12:
            return m.group(0)
        return f"{_MONTHS_EN_SHORT[mo - 1]} {d:02d}, {y:04d}"

    return re.sub(r"(\d{4})[-./](\d{1,2})[-./](\d{1,2})", _sub, out)


def _derive_incident(
    ref_raw: str | None,
    cedant_code: str | None,
) -> tuple[str | None, int | None]:
    """Derive (incident_id, claim_seq) from a row's ref_no.

    ``incident_id`` is the ref body with the revision suffix stripped,
    so successive billings of the same underlying claim
    (``2020-1021024610 001`` / ``002`` / ``003``) collapse to a single
    identity. ``claim_seq`` is the integer revision (1, 2, 3 …) or
    ``None`` when the source ref carries no revision component.

    Returns ``(None, None)`` for empty input, unsupported cedants (we
    only have a verified parser for KB today), or sentinel values
    (``"various"`` etc.). Per-cedant parsers for KR/HW/SS/DB/HM can be
    plugged in here as they land via the playbook system.
    """
    if not ref_raw:
        return None, None
    s = str(ref_raw).strip()
    if not s:
        return None, None
    if (cedant_code or "").upper() != "KB":
        return None, None
    try:
        parsed = parse_kb_ref_no(s)
    except RefNoParseError:
        return None, None
    if parsed.is_sentinel:
        return None, None
    return parsed.body, parsed.revision

router = APIRouter(prefix="/api/tools/soc", tags=["soc_stream"])
logger = logging.getLogger(__name__)


def _parse_json(raw: str):
    """Parse JSON, unwrapping a list response to its first element."""
    return _parse_json_raw(raw, unwrap_list=True)


def _sse(event: str, data: dict) -> str:
    return f"event: {event}\ndata: {json.dumps(data, ensure_ascii=False, default=str)}\n\n"


@router.post("/stream")
async def stream_soc(
    files: list[UploadFile] = File(default=[]),
    text: str | None = Form(None),
):
    # NOTE: we intentionally do NOT use Depends(get_db) here. FastAPI dependency
    # cleanup for StreamingResponse is racy — the session can get GC'd before
    # the generator finishes, leaving the pooled connection dangling (visible
    # as "garbage collector is trying to clean up non-checked-in connection"
    # warnings). Instead, manage the session's lifetime explicitly inside the
    # generator with `async with async_session() as db` so it is guaranteed to
    # close when the stream ends or is cancelled.

    # Read files BEFORE entering the generator (UploadFile closes after response starts)
    saved_files = []
    for file in files:
        data = await file.read()
        suffix = Path(file.filename or "").suffix.lower()
        f = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
        f.write(data)
        f.flush()
        f.close()
        saved_files.append((f.name, file.filename, suffix))

    saved_text = text

    async def generate() -> AsyncGenerator[str, None]:
        import time as _t
        # Prime the stream so browsers/proxies stop buffering and start
        # delivering events immediately. SSE comment lines (`: ...`) are ignored
        # by the EventSource parser but force the TCP connection to flush.
        yield ": stream-open\n\n"
        t0 = _t.time()
        tmp_paths = saved_files
        # Pre-generate a case ID and emit it immediately so the frontend can
        # swap to a permalink URL even before the AI pipeline finishes.
        case_id = _uuid.uuid4()
        # Open the DB session for the lifetime of this generator. `async with`
        # guarantees close() runs even if the client disconnects and the
        # generator is cancelled.
        async with async_session() as db:
            # Phase 3c: initialise KB deterministic-parser context.
            # Inactive (no-op) when KB_PARSER_MODE=OFF.
            kb_ctx = init_kb_integration(db_session=db)
            case = SlipCase(
                id=case_id,
                status="processing",
                input_text=saved_text,
                input_files=[
                    {"name": fname, "suffix": suffix}
                    for _, fname, suffix in tmp_paths
                ],
            )
            db.add(case)
            await db.commit()
            yield _sse("case", {"id": str(case_id)})
            yield _sse("step", {"phase": "upload", "message": f"📂 파일 {len(tmp_paths)}개 수신 완료"})
            for _, fname, suffix in tmp_paths:
                yield _sse("step", {"phase": "upload", "message": f"  · {fname} ({suffix or '?'})"})
    
            # STEP 1: Upload to Gemini + extract
            yield _sse("step", {"phase": "extract", "message": "🔧 Gemini API 초기화"})
            genai.configure(api_key=settings.gemini_api_key)
            model = genai.GenerativeModel("gemini-3.1-flash-lite")
            yield _sse("step", {"phase": "extract", "message": "  모델: gemini-3.1-flash-lite"})
    
            yield _sse("step", {"phase": "extract", "message": f"📤 파일 {len(tmp_paths)}개 업로드 시작"})
            parts = []
            uploaded_ct = 0
            failed_ct = 0
            skipped_ct = 0
            # Gemini File API supports PDFs/images/text/audio/video. xlsx/doc(x)/etc. can't be processed,
            # so we parse supported types locally and upload the result as text.
            _MIME_BY_SUFFIX = {
                ".pdf": "application/pdf",
                ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
                ".gif": "image/gif", ".webp": "image/webp",
                ".txt": "text/plain", ".csv": "text/csv", ".html": "text/html", ".htm": "text/html",
                ".md": "text/markdown",
                ".mp3": "audio/mpeg", ".wav": "audio/wav", ".m4a": "audio/mp4",
                ".mp4": "video/mp4", ".mov": "video/quicktime",
            }
            for i, (tmp_path, filename, suffix) in enumerate(tmp_paths, 1):
                yield _sse("step", {"phase": "extract", "message": f"  [{i}/{len(tmp_paths)}] {filename} 업로드 중..."})
                try:
                    # xlsx / xls → read locally, attach as text
                    if suffix in (".xlsx", ".xls"):
                        # Phase 3c: KB deterministic parser entry. Returns
                        # an LLM-context string only in ON mode; SHADOW /
                        # OFF / non-KB inputs return None and we fall
                        # through to the legacy text-dump path.
                        kb_context_str = try_parse_kb_xlsx(kb_ctx, tmp_path, filename)
                        if kb_context_str:
                            parts.append(kb_context_str)
                        try:
                            from openpyxl import load_workbook
                            wb = load_workbook(tmp_path, data_only=True, read_only=True)
                            lines = [f"[Excel: {filename}]"]
                            for sh in wb.sheetnames[:5]:
                                ws = wb[sh]
                                lines.append(f"\n## Sheet: {sh}")
                                for row in ws.iter_rows(values_only=True, max_row=200):
                                    row_txt = "\t".join("" if c is None else str(c) for c in row)
                                    if row_txt.strip(): lines.append(row_txt)
                            parts.append("\n".join(lines)[:20000])
                            uploaded_ct += 1
                            yield _sse("step", {"phase": "extract", "message": f"  ✓ [{i}/{len(tmp_paths)}] 엑셀 내용 추출: {filename}"})
                        except Exception as ee:
                            skipped_ct += 1
                            yield _sse("step", {"phase": "extract", "message": f"  ⊘ [{i}/{len(tmp_paths)}] 엑셀 파싱 스킵: {filename} — {str(ee)[:80]}"})
                        continue
                    # msg → extract .msg via extract_msg if available; else skip
                    if suffix == ".msg":
                        try:
                            import extract_msg
                            msg = extract_msg.Message(tmp_path)
                            body = f"[Email: {filename}]\nSubject: {msg.subject or ''}\nFrom: {msg.sender or ''}\nTo: {msg.to or ''}\nDate: {msg.date or ''}\n\n{msg.body or ''}"
                            parts.append(body[:20000])
                            uploaded_ct += 1
                            yield _sse("step", {"phase": "extract", "message": f"  ✓ [{i}/{len(tmp_paths)}] 메일 본문 추출: {filename}"})
                        except Exception as ee:
                            skipped_ct += 1
                            yield _sse("step", {"phase": "extract", "message": f"  ⊘ [{i}/{len(tmp_paths)}] 메일 파싱 스킵: {filename} — {str(ee)[:80]}"})
                        continue
    
                    mime = _MIME_BY_SUFFIX.get(suffix)
                    if mime:
                        uploaded = await asyncio.to_thread(genai.upload_file, tmp_path, mime_type=mime)
                    else:
                        uploaded = await asyncio.to_thread(genai.upload_file, tmp_path)
                    # Wait for the file to reach ACTIVE state — otherwise Gemini returns
                    # "The document has no pages" when it's still PROCESSING.
                    import time as _time
                    wait_start = _time.time()
                    while getattr(uploaded, "state", None) and uploaded.state.name == "PROCESSING":
                        if _time.time() - wait_start > 30:
                            break
                        await asyncio.sleep(0.5)
                        uploaded = await asyncio.to_thread(genai.get_file, uploaded.name)
                    if getattr(uploaded, "state", None) and uploaded.state.name == "FAILED":
                        failed_ct += 1
                        yield _sse("step", {"phase": "extract", "message": f"  ✗ [{i}/{len(tmp_paths)}] 처리 실패: {filename}"})
                        continue
                    parts.append(uploaded)
                    parts.append(f"[File: {filename}]")
                    uploaded_ct += 1
                    yield _sse("step", {"phase": "extract", "message": f"  ✓ [{i}/{len(tmp_paths)}] 업로드 완료: {filename}"})
                except Exception as e:
                    failed_ct += 1
                    yield _sse("step", {"phase": "extract", "message": f"  ✗ [{i}/{len(tmp_paths)}] 업로드 실패: {filename} — {str(e)[:100]}"})
    
            yield _sse("step", {"phase": "extract", "message": f"📊 업로드 결과: 성공 {uploaded_ct} / 실패 {failed_ct}" + (f" / 스킵 {skipped_ct}" if skipped_ct else "")})
    
            if saved_text and saved_text.strip():
                parts.append(f"[Text]:\n{text[:3000]}")
                yield _sse("step", {"phase": "extract", "message": f"📝 추가 텍스트 {len(text)}자 첨부"})
    
            yield _sse("step", {"phase": "extract", "message": "🤖 AI에게 클레임 정보 추출 요청 중... (PASS 1/2)"})
    
            parts.append(EXTRACT_CLAIM_BORDEREAU)
    
            t_ai1 = _t.time()
            try:
                resp = await asyncio.to_thread(model.generate_content, parts)
                ai_result = _parse_json(resp.text)
                yield _sse("step", {"phase": "extract", "message": f"✓ AI 응답 수신 ({_t.time() - t_ai1:.1f}초)"})
            except Exception as e:
                yield _sse("error", {"message": f"AI 추출 실패: {e}"})
                case.status = "error"
                case.error_message = f"AI 추출 실패: {e}"
                await db.merge(case)
                await db.commit()
                yield _sse("done", {})
                return

            if not ai_result or not isinstance(ai_result, dict):
                yield _sse("error", {"message": "AI 응답이 유효하지 않습니다"})
                case.status = "error"
                case.error_message = "AI 응답이 유효하지 않습니다"
                await db.merge(case)
                await db.commit()
                yield _sse("done", {})
                return

            # Phase 3c: cross-check the LLM extraction against the
            # deterministic parser output. Audit-only — never mutates
            # ai_result. Surface fail / warning severities as SSE step
            # events so the operator sees them in the live stream.
            for cc in run_kb_cross_checks(kb_ctx, ai_result):
                if cc.severity == SEVERITY_FAIL:
                    yield _sse("step", {"phase": "extract", "message": f"⚠ KB cross-check FAIL: rows {cc.deterministic_row_count} vs {cc.llm_row_count}, {len(cc.mismatch_details)} mismatches"})
                elif cc.severity == SEVERITY_WARNING:
                    yield _sse("step", {"phase": "extract", "message": f"⚠ KB cross-check warning: {len(cc.mismatch_details)} mismatch(es)"})

            raw_claims = ai_result.get("claims") or []
            norm_claims = []
            _to_num = _to_num_util
            for rc in raw_claims if isinstance(raw_claims, list) else []:
                if not isinstance(rc, dict):
                    continue
                row_cedant = (
                    rc.get("cedant") or ai_result.get("cedant_code") or ""
                )
                row_ref = (
                    rc.get("cedant_ref_no") or ai_result.get("ref_no") or ""
                )
                incident_id, claim_seq = _derive_incident(row_ref, row_cedant)
                norm_claims.append({
                    "insured": rc.get("insured") or ai_result.get("account_name") or "",
                    # Default Policy Period to KB-style English
                    # (``Mmm DD, YYYY`` zero-padded). Idempotent on
                    # already-English input; ISO dates get rewritten.
                    "policy_period": _to_english_date_range(
                        rc.get("policy_period") or ai_result.get("policy_period") or ""
                    ),
                    "uy": str(rc.get("uy") or ""),
                    "facility": rc.get("facility") or rc.get("cedant") or ai_result.get("cedant_code") or "",
                    "dol": rc.get("dol") or ai_result.get("dol") or "",
                    "doc": rc.get("doc") or "",
                    "currency": rc.get("currency") or ai_result.get("currency") or "KRW",
                    "claim_amount_100": _to_num(rc.get("claim_amount_100")),
                    "expense_100": _to_num(rc.get("expense_100")),
                    "cedant": row_cedant,
                    "cedant_ref_no": row_ref,
                    # Same-incident identity: rows that share ``incident_id``
                    # are successive billings (``001``/``002`` …) of one
                    # claim. Customers identify same-incident by
                    # (incident_id, dol). ``claim_seq`` is the revision int.
                    "incident_id": incident_id,
                    "claim_seq": claim_seq,
                })

            # Cedants frequently send PDFs with the reinsured/reinsurer name
            # in ALL-CAPS letterhead style (e.g. "KOREAN REINSURANCE COMPANY").
            # Canonicalize against the Excel-derived code table so the slip
            # uses the Title-Case form humans use when typing slips by hand.
            # Reinsured = the cedant company being reinsured. Per the
            # business team, slip text should always use the canonical
            # name from CODE_관리.xlsx — never the raw AI-extracted
            # string (which can be all-caps with trailing dots etc.).
            # Resolve in priority: AI's reinsured → cedant_code lookup
            # → AI's reinsured as-is.
            _ai_reinsured = (ai_result.get("reinsured") or "").strip()
            _cedant_code = (ai_result.get("cedant_code") or "").strip()
            _canon_reinsured = canonicalize_company_name(_ai_reinsured) if _ai_reinsured else ""
            if not _canon_reinsured and _cedant_code:
                _canon_reinsured = canonicalize_company_name(_cedant_code)
            if not _canon_reinsured:
                _canon_reinsured = _ai_reinsured

            extracted = {
                "account_name": ai_result.get("account_name"),
                "claimant": ai_result.get("claimant") or "",
                "policy_no": ai_result.get("policy_no") or "",
                "reinsured": _canon_reinsured,
                "cedant": ai_result.get("cedant_code"),
                "line": ai_result.get("line"),
                "ref_no": ai_result.get("ref_no"),
                "dol": ai_result.get("dol"),
                "doc_date": ai_result.get("doc_date") or "",
                "title_subline": ai_result.get("title_subline") or "",
                "currency": ai_result.get("currency"),
                "total_amount": ai_result.get("total_amount_100"),
                "expenses_reserve": _to_num_util(ai_result.get("expenses_reserve_100")),
                "location_of_loss": ai_result.get("location_of_loss"),
                "nature_of_loss": ai_result.get("nature_of_loss"),
                "particulars": ai_result.get("particulars"),
                "policy_period": ai_result.get("policy_period"),
                "description": ai_result.get("description"),
                "claims": norm_claims,
            }
            try:
                extracted["total_amount"] = float(str(extracted.get("total_amount") or 0).replace(",", ""))
            except:
                extracted["total_amount"] = 0
            if (not extracted["total_amount"]) and norm_claims:
                extracted["total_amount"] = sum(c["claim_amount_100"] + c["expense_100"] for c in norm_claims)

            # Reconcile: if AI reported a document Total that disagrees with sum of rows,
            # try to auto-flip a row's sign (common failure mode: AI drops the "-" on a
            # reversal/adjustment row in a Korean bordereau).
            if norm_claims and extracted["total_amount"]:
                reported = float(extracted["total_amount"])
                row_sum = sum(c["claim_amount_100"] + c["expense_100"] for c in norm_claims)
                diff = row_sum - reported
                if abs(diff) > 1:
                    flipped = False
                    for c in norm_claims:
                        row_val = c["claim_amount_100"] + c["expense_100"]
                        # If flipping this row's signs closes the gap, do it.
                        if row_val > 0 and abs(2 * row_val - diff) < 1:
                            c["claim_amount_100"] = -c["claim_amount_100"]
                            c["expense_100"] = -c["expense_100"]
                            flipped = True
                            yield _sse("step", {"phase": "extract", "message": f"⚠ 행 합계 불일치 감지 — Ref {c.get('cedant_ref_no','?')} 의 부호를 음수로 교정 (환수/조정 row)"})
                            break
                    if not flipped:
                        yield _sse("step", {"phase": "extract", "message": f"⚠ 행 합계({row_sum:,.0f}) ≠ 문서 Total({reported:,.0f}) — 수동 확인 필요"})
    
            yield _sse("extracted", {"data": extracted})
            yield _sse("step", {"phase": "extract", "message": f"📋 피보험자: {extracted.get('account_name') or '-'}"})
            yield _sse("step", {"phase": "extract", "message": f"💰 금액: {extracted.get('currency') or '-'} {extracted.get('total_amount') or 0:,.0f}"})
            if ai_result.get("file_reinsurer_name"):
                yield _sse("step", {"phase": "extract", "message": f"🏷 파일 내 재보험사 힌트: {ai_result['file_reinsurer_name']} / {float(ai_result.get('file_share') or 0)*100:.2f}% / {float(ai_result.get('file_amount') or 0):,.0f}"})
    
            file_share = float(str(ai_result.get("file_share") or 0).replace(",", ""))
            file_amount = float(str(ai_result.get("file_amount") or 0).replace(",", ""))
            insured_name = extracted.get("account_name", "")
            # Some cedants (notably KB) mask the Assured name in the SOC body
            # ("신****인") but include the unmasked entity in a separate
            # "Claimant:" line. ILIKE search on the masked form will never
            # match the canonical DB entry, so fall back to claimant for
            # name-based lookups when masking asterisks are present.
            claimant_name = extracted.get("claimant", "") or ""
            if "*" in (insured_name or "") and claimant_name:
                yield _sse("step", {"phase": "extract", "message": f"  ℹ Insured 마스킹 감지 ({insured_name}) → 검색은 Claimant 사용: {claimant_name}"})
                insured_name = claimant_name
            ref_no = extracted.get("ref_no", "")
    
            # Try to extract ref_no from filenames (more reliable than AI for scanned docs).
            # Prefer the structured cedant-ref shape (YYYYMMDD-NNNNN-NN, optional "(N)")
            # over a bare digit run — otherwise a date embedded in the filename
            # (e.g. SamSungFire_oz_20260321_*.pdf) clobbers the real cedant ref the
            # AI already pulled from the document body.
            structured_re = re.compile(r'(\d{8}-\d{4,6}-\d{2}(?:\s*\(\d+\))?)')
            ai_ref_is_structured = bool(re.search(r'\d-\d', ref_no or ""))
            for _, fname, _ in tmp_paths:
                fname_s = fname or ""
                m = structured_re.search(fname_s)
                if m:
                    fname_ref = m.group(1).strip()
                    if fname_ref != ref_no:
                        ref_no = fname_ref
                        extracted["ref_no"] = ref_no
                        yield _sse("step", {"phase": "extract", "message": f"파일명에서 Ref 추출: {ref_no}"})
                    break
                if ai_ref_is_structured:
                    # AI already has a structured ref from the document body —
                    # don't replace it with a bare date-like digit run.
                    continue
                ref_match = re.search(r'(\d{8,})', fname_s)
                if ref_match:
                    fname_ref = ref_match.group(1)
                    if fname_ref != ref_no:
                        ref_no = fname_ref
                        extracted["ref_no"] = ref_no
                        yield _sse("step", {"phase": "extract", "message": f"파일명에서 Ref 추출 (fallback): {ref_no}"})
    
            # STEP 2: Search past claims
            yield _sse("step", {"phase": "claims", "message": "🔎 STEP 2: 과거 클레임 DB 검색"})
    
            past_claims_data: list[dict] = []
            claims_method = "no match"
            if ref_no:
                # Strip trailing order/suffix that varies per claim line:
                #   "2013-0000539637-155" (dash) → base "2013-0000539637"
                #   "2013-0000539637 155" (space) → base "2013-0000539637"
                #   "C2025018772029-1-7" (multi-dash, single-digit) → base "C2025018772029-1"
                base_ref = re.sub(r'[\s\-]+\d{1,4}$', '', ref_no).strip()
                yield _sse("step", {"phase": "claims", "message": f"  · Ref 기반: SELECT FROM claims WHERE ref_no ILIKE '%{base_ref}%'"})
                found = (await db.execute(select(Claim).where(Claim.ref_no.ilike(f"%{base_ref}%")).limit(20))).scalars().all()
                if found:
                    claims_method = f"ref ILIKE %{base_ref}%"
                    past_claims_data = [_claim_to_dict(cl) for cl in found]
                    yield _sse("step", {"phase": "claims", "message": f"  ✓ {len(found)}건 발견"})
                else:
                    yield _sse("step", {"phase": "claims", "message": "  · Ref 매칭 없음"})
                if found and not insured_name:
                    insured_name = found[0].account_name or ""
                    extracted["account_name"] = insured_name
                    yield _sse("step", {"phase": "claims", "message": f"  · DB에서 피보험자 확인: {insured_name}"})
                    yield _sse("extracted", {"data": extracted})
    
            if not past_claims_data and insured_name:
                # Try progressively broader keys: full name → corp-suffix-stripped → 15-char prefix
                search_keys: list[str] = []
                for k in (insured_name, strip_corp_suffix(insured_name), insured_name[:15]):
                    k = (k or "").strip()
                    if k and k not in search_keys:
                        search_keys.append(k)
                for key in search_keys:
                    yield _sse("step", {"phase": "claims", "message": f"  · 피보험자 기반: SELECT FROM claims WHERE account_name ILIKE '%{key}%'"})
                    found = (await db.execute(select(Claim).where(Claim.account_name.ilike(f"%{key}%")).limit(20))).scalars().all()
                    if found:
                        claims_method = f"ILIKE %{key}%"
                        past_claims_data = [_claim_to_dict(cl) for cl in found]
                        yield _sse("step", {"phase": "claims", "message": f"  ✓ {len(found)}건 발견"})
                        break
                    yield _sse("step", {"phase": "claims", "message": "  · 매칭 없음"})
    
            past_claims = [_fmt_claim(d) for d in past_claims_data]
            yield _sse("step", {"phase": "claims", "message": f"📦 과거 클레임 최종 {len(past_claims)}건"})
            if past_claims_data:
                distinct_ri = sorted({d.get("reinsurer") for d in past_claims_data if d.get("reinsurer")})
                if distinct_ri:
                    yield _sse("step", {"phase": "claims", "message": f"  · 참여 재보험사: {', '.join(distinct_ri[:8])}"})

            # Multi-cedant / multi-reinsurer fan-out detection runs AFTER
            # the contracts search step (it needs contracts_data as a
            # fallback source when past_claims is empty). Initialised here
            # so it stays in scope through the SOCs build later.
            fanout_groups: dict[tuple[str, str], dict] = {}
            distinct_cedants: list[str] = []
            multi_cedant = False
    
            # STEP 3: Search contracts
            yield _sse("step", {"phase": "contracts", "message": "🔎 STEP 3: 계약(Contract) DB 검색"})
    
            contracts_data: list[dict] = []
            contracts_method = "no match"
            if insured_name:
                search_keys: list[str] = []
                for k in (insured_name, strip_corp_suffix(insured_name), insured_name[:15]):
                    k = (k or "").strip()
                    if k and k not in search_keys:
                        search_keys.append(k)
                for key in search_keys:
                    yield _sse("step", {"phase": "contracts", "message": f"  · 피보험자 기반: SELECT FROM contracts WHERE assured ILIKE '%{key}%' AND share > 0"})
                    found = (await db.execute(select(Contract).where(Contract.assured.ilike(f"%{key}%"), Contract.share > 0).limit(20))).scalars().all()
                    if found:
                        contracts_method = f"ILIKE %{key}%"
                        contracts_data = [_contract_to_dict(c) for c in found]
                        yield _sse("step", {"phase": "contracts", "message": f"  ✓ {len(found)}건 발견"})
                        break
                    yield _sse("step", {"phase": "contracts", "message": "  · 매칭 없음"})
    
            matching_contracts = [_fmt_contract(d) for d in contracts_data]
            yield _sse("step", {"phase": "contracts", "message": f"📦 계약 최종 {len(matching_contracts)}건"})
            if contracts_data:
                distinct_ri = sorted({d.get("reinsurer") for d in contracts_data if d.get("reinsurer")})
                if distinct_ri:
                    yield _sse("step", {"phase": "contracts", "message": f"  · 참여 재보험사: {', '.join(distinct_ri[:8])}"})

            # If the source PDF masked the Insured ("신****인"), now that we
            # have past_claims and contracts populated, resolve to the real
            # name and re-emit `extracted` so the frontend swaps the masked
            # form for the canonical one.
            if "*" in (extracted.get("account_name") or ""):
                resolved_name, resolved_src = resolve_insured_name(
                    extracted.get("account_name") or "",
                    claimant=claimant_name,
                    past_claims=past_claims_data,
                    contracts=contracts_data,
                    policy=None,
                )
                if resolved_name and "*" not in resolved_name:
                    extracted["account_name"] = resolved_name
                    insured_name = resolved_name
                    # Propagate the resolved name into every bordereau
                    # row whose ``insured`` is still masked. AI extracts
                    # the per-row Insured independently (PDF often
                    # masks every row), so we patch them all here once
                    # we have the canonical form from past_claims /
                    # contracts. Without this, the rendered slip table
                    # keeps showing ``LX ****************`` even though
                    # the slip header is unmasked.
                    patched = 0
                    for row in norm_claims:
                        if "*" in (row.get("insured") or ""):
                            row["insured"] = resolved_name
                            patched += 1
                    yield _sse("step", {
                        "phase": "extract",
                        "message": f"  ✓ 마스킹 해제: '{resolved_name}' (출처: {resolved_src})"
                                   + (f" · {patched}개 행에 적용" if patched else ""),
                    })
                    yield _sse("extracted", {"data": extracted})
                else:
                    yield _sse("step", {
                        "phase": "extract",
                        "message": "  ⚠ 마스킹 해제 실패 — 매칭 데이터에서 후보를 찾지 못함",
                    })

            # ─── (cedant, reinsurer) fan-out detection ────────────
            # Generates one slip per pair when the data shows the new
            # loss is shared. Past claims are preferred (actual
            # settlements). When they're empty/single, fall back to
            # active contracts whose policy period covers the dol —
            # contracts encode the same plan past claims would realise.
            for c in past_claims_data:
                ced = (c.get("cedant") or "").strip()
                ri = (c.get("reinsurer") or "").strip()
                if not ced or not ri or is_broker_name(ri):
                    continue
                sh = float(c.get("share") or 0)
                amt = float(c.get("krw_amount") or 0)
                k = (ced, ri)
                g = fanout_groups.setdefault(
                    k,
                    {"sum_share": 0.0, "count": 0, "sum_amt": 0.0, "source": "CLAIMS"},
                )
                g["sum_share"] += sh
                g["count"] += 1
                g["sum_amt"] += amt

            from datetime import date as _date
            def _covers_dol(pf, pt, dol_str):
                if not dol_str or not pf or not pt:
                    return True
                try:
                    d = _date.fromisoformat(dol_str.replace("/", "-")[:10])
                    return _date.fromisoformat(pf[:10]) <= d <= _date.fromisoformat(pt[:10])
                except Exception:
                    return True

            if len(fanout_groups) < 2:
                dol_str = extracted.get("dol") or ""
                for c in contracts_data:
                    if not _covers_dol(c.get("period_from"), c.get("period_to"), dol_str):
                        continue
                    ced = (c.get("cedant") or "").strip()
                    ri = (c.get("reinsurer") or "").strip()
                    sh = float(c.get("share") or 0)
                    if not ced or not ri or sh <= 0 or is_broker_name(ri):
                        continue
                    k = (ced, ri)
                    if k in fanout_groups and fanout_groups[k]["source"] == "CLAIMS":
                        continue
                    g = fanout_groups.setdefault(
                        k,
                        {"sum_share": 0.0, "count": 0, "sum_amt": 0.0, "source": "CONTRACTS"},
                    )
                    g["sum_share"] += sh
                    g["count"] += 1
                    g["sum_amt"] += float(c.get("ri_prem") or 0)
            distinct_cedants = sorted({k[0] for k in fanout_groups})
            multi_cedant = len(fanout_groups) >= 2
            if multi_cedant:
                source_mix = sorted({g["source"] for g in fanout_groups.values()})
                pair_summary = ", ".join(f"{c}/{r}" for c, r in fanout_groups.keys())
                yield _sse("step", {"phase": "contracts", "message": f"⚠ Co-insurance 감지: {len(fanout_groups)}개 (cedant, reinsurer) 페어 [{'/'.join(source_mix)}] → fanout 모드"})
                yield _sse("step", {"phase": "contracts", "message": f"  · 페어: {pair_summary}"})
    
            # STEP 4: Find policy
            yield _sse("step", {"phase": "policy", "message": "🔎 STEP 4: 폴리시 매칭 (점수 기반)"})
    
            # Simple policy match
            policy_data = None
            if insured_name:
                def _norm(s): return re.sub(r'\([^)]*\)', '', (s or "").strip().lower().replace("co.,ltd", "").replace("inc.", "").replace("corp.", "").replace("ltd.", "")).strip()
                norm_cedant = normalize_cedant(extracted.get("cedant") or "")
    
                policies = (await db.execute(select(Policy))).scalars().all()
                valid_cns = set((await db.execute(
                    select(Contract.cover_note_no).where(Contract.share > 0).distinct()
                )).scalars().all())
                norm = _norm(insured_name)
                best, best_score = None, 0
                for p in policies:
                    if p.cover_note_no not in valid_cns: continue
                    score = 0
                    pn = _norm(p.assured)
                    if pn and norm:
                        if pn == norm: score += 40
                        elif pn in norm or norm in pn: score += 30
                    if p.cedant and (p.cedant == norm_cedant): score += 20
                    if score > best_score: best_score = score; best = p
                if best and best_score >= 30:
                    policy_data = {"id": best.id, "cover_note_no": best.cover_note_no, "assured": best.assured, "cedant": best.cedant, "score": best_score}
    
            if policy_data:
                yield _sse("step", {"phase": "policy", "message": f"  ✓ 매칭: {policy_data['cover_note_no']} / {policy_data['assured']} (점수: {policy_data['score']})"})
                yield _sse("policy", {"data": policy_data})
            else:
                yield _sse("step", {"phase": "policy", "message": "  · 매칭되는 폴리시 없음"})
    
            # STEP 5: AI reasoning for reinsurer selection — skipped in
            # multi-cedant mode where past data drives a deterministic
            # per-cedant fan-out instead.
            reasoning_result = None
            if multi_cedant:
                yield _sse("step", {"phase": "reasoning", "message": "🧠 STEP 5: AI reasoning 스킵 (다중 cedant — 과거 데이터 기반 fanout)"})
            else:
                yield _sse("step", {"phase": "reasoning", "message": "🧠 STEP 5: AI에게 재보험사 배분 판단 요청 (PASS 2/2)"})
                yield _sse("step", {"phase": "reasoning", "message": f"  · 증거: 파일={1 if ai_result.get('file_reinsurer_name') else 0}, 과거클레임={len(past_claims)}, 계약={len(matching_contracts)}, 폴리시={'있음' if policy_data else '없음'}"})

                reasoning_prompt = REINSURER_REASONING_SHORT.format(
                    insured_name=insured_name,
                    line=extracted.get("line"),
                    ref_no=ref_no,
                    file_ri_name=ai_result.get("file_reinsurer_name", ""),
                    file_share_pct=file_share * 100,
                    file_amount=file_amount,
                    past_claims_block=chr(10).join(past_claims[:15]) if past_claims else "(없음)",
                    contracts_block=chr(10).join(matching_contracts[:15]) if matching_contracts else "(없음)",
                )

                t_ai2 = _t.time()
                try:
                    resp2 = await asyncio.to_thread(model.generate_content, reasoning_prompt)
                    reasoning_result = _parse_json(resp2.text)
                    yield _sse("step", {"phase": "reasoning", "message": f"  ✓ AI 추론 완료 ({_t.time() - t_ai2:.1f}초)"})
                except Exception as e:
                    yield _sse("step", {"phase": "reasoning", "message": f"  ✗ 추론 실패: {e}"})
                    reasoning_result = None
    
            # Stream thinking steps
            thinking = []
            if reasoning_result and isinstance(reasoning_result, dict):
                thinking = reasoning_result.get("thinking", [])
                if thinking:
                    yield _sse("step", {"phase": "reasoning", "message": f"  · CoT {len(thinking)}단계 생성됨"})
                for i, step in enumerate(thinking):
                    yield _sse("thinking", {"step": i + 1, "content": step})
    
            # Build SOCs
            ccy = extracted.get("currency", "KRW")
            total = float(extracted.get("total_amount") or 0)
            socs = []

            if multi_cedant:
                # ─── (cedant, reinsurer) pair fan-out ────────────
                # Co-insurance / multi-reinsurer scenario: one PLA per
                # detected (cedant, reinsurer) pair. Share = average
                # historical share for that pair; amount = total × share.
                # Source is "CLAIMS" when the pair came from past
                # settlements, "CONTRACTS" when it came from the active
                # contracts fallback. AI reasoning is skipped because the
                # answer is fully data-driven.
                yield _sse("step", {"phase": "reasoning", "message": "🔀 Co-insurance fanout: AI reasoning 우회, 데이터 기반 페어 생성"})
                for (ced, ri_name), g in fanout_groups.items():
                    avg_share = g["sum_share"] / max(g["count"], 1)
                    amount = total * avg_share
                    src = g["source"]
                    src_label = "과거 클레임" if src == "CLAIMS" else "활성 계약"
                    socs.append({
                        "reinsurer": resolve_full_name(ri_name),
                        "share": avg_share,
                        "share_pct": f"{avg_share * 100:.2f}%",
                        "amount": amount,
                        "amount_fmt": _fmt_amount(ccy, amount),
                        # Cedant code preserved for backward compat;
                        # cedant_full_name is the canonical insurer name
                        # the slip should display in its "Reinsured" row
                        # — different per fanout slip when the underlying
                        # cedants differ.
                        "cedant": ced,
                        "cedant_full_name": canonicalize_company_name(ced) or ced,
                        "confidence": f"FANOUT_{src}",
                        "reasoning": (
                            f"({ced}, {ri_name}) 페어 — {src_label} {g['count']}건 기반, "
                            f"평균 share {avg_share * 100:.2f}%."
                        ),
                        "claims": [dict(c) for c in norm_claims],
                    })
                    yield _sse("step", {"phase": "reasoning", "message": f"  ✓ {ced} / {ri_name} ({avg_share*100:.2f}%) [{src}]"})
            elif reasoning_result and isinstance(reasoning_result, dict):
                for ri in (reasoning_result.get("reinsurers") or []):
                    name = ri.get("name", "")
                    if not name or is_broker_name(name): continue
                    share = float(str(ri.get("share") or 0).replace(",", ""))
                    if file_amount and len(reasoning_result.get("reinsurers", [])) == 1:
                        amount = file_amount
                    else:
                        amount = total * share
                    socs.append({
                        "reinsurer": resolve_full_name(name), "share": share,
                        "share_pct": f"{share * 100:.2f}%",
                        "amount": amount, "amount_fmt": _fmt_amount(ccy, amount),
                        "cedant": extracted.get("cedant", ""),
                        "confidence": ri.get("source", "AI"),
                        "reasoning": ri.get("reasoning", ""),
                        "claims": [dict(c) for c in norm_claims],
                    })
    
            # Fallback to policy if no reasoning result
            if not socs and policy_data:
                yield _sse("step", {"phase": "fallback", "message": "폴리시 기반 폴백 사용 중..."})
                contracts = (await db.execute(select(Contract).where(Contract.cover_note_no == policy_data["cover_note_no"], Contract.share > 0))).scalars().all()
                for c in contracts:
                    share = float(c.share or 0)
                    amount = total * share
                    socs.append({
                        "reinsurer": resolve_full_name(c.reinsurer), "share": share,
                        "share_pct": f"{share * 100:.2f}%",
                        "amount": amount, "amount_fmt": _fmt_amount(ccy, amount),
                        "cedant": c.cedant or "", "confidence": "POLICY", "reasoning": f"계약 {policy_data['cover_note_no']} 기반",
                        "claims": [dict(c) for c in norm_claims],
                    })
    
            # Tag each SOC with foreign-reinsurer classification so the
            # frontend can render the right language variant of the slip.
            for s in socs:
                s["is_foreign"] = is_foreign_reinsurer(s.get("reinsurer", ""))

            # If any SOC is for a foreign reinsurer, translate the Korean
            # text fields once into English and attach to extracted.text_en.
            # The frontend picks Korean (extracted.*) for domestic slips and
            # English (extracted.text_en.*) for foreign slips. We translate
            # only once because the source text is the same for every slip.
            any_foreign = any(s["is_foreign"] for s in socs)
            if any_foreign:
                # Insured: prefer the English form already present in
                # contracts/policy (e.g. "Korea South-East Power Co., Ltd.
                # (한국남동발전)" → "Korea South-East Power Co., Ltd."). When
                # no DB-backed English exists we fall through to LLM
                # translation below.
                en_insured, en_src = pick_english_insured_name(
                    extracted.get("account_name") or "",
                    contracts=contracts_data,
                    policy=policy_data,
                )
                text_en_seed: dict = {}
                if en_insured:
                    text_en_seed["account_name"] = en_insured
                    yield _sse("step", {
                        "phase": "translate",
                        "message": f"🌐 Insured 영문화 (DB): '{en_insured}' (출처: {en_src})",
                    })

                yield _sse("step", {"phase": "translate", "message": "🌐 외국 재보험사 슬립 — 한→영 번역 시작"})
                tr_prompt = SLIP_TRANSLATE_TO_ENGLISH.format(
                    account_name=extracted.get("account_name") or "",
                    line=extracted.get("line") or "",
                    location_of_loss=extracted.get("location_of_loss") or "",
                    nature_of_loss=extracted.get("nature_of_loss") or "",
                    particulars=extracted.get("particulars") or extracted.get("description") or "",
                    remarks=extracted.get("remarks") or "",
                    description=extracted.get("description") or "",
                )
                t_tr = _t.time()
                try:
                    resp_tr = await asyncio.to_thread(model.generate_content, tr_prompt)
                    text_en = _parse_json(resp_tr.text)
                    if isinstance(text_en, dict):
                        # DB-resolved English wins over LLM output for the
                        # company name (LLM can hallucinate proper nouns).
                        if "account_name" in text_en_seed:
                            text_en["account_name"] = text_en_seed["account_name"]
                        extracted["text_en"] = text_en
                        yield _sse("step", {"phase": "translate", "message": f"  ✓ 번역 완료 ({_t.time() - t_tr:.1f}초)"})
                    else:
                        if text_en_seed:
                            extracted["text_en"] = text_en_seed
                        yield _sse("step", {"phase": "translate", "message": "  ✗ 번역 응답 형식 오류 — 한글 fallback"})
                except Exception as e:
                    if text_en_seed:
                        extracted["text_en"] = text_en_seed
                    yield _sse("step", {"phase": "translate", "message": f"  ✗ 번역 실패: {e} — 한글 fallback"})

            # Summary
            yield _sse("step", {"phase": "done", "message": f"✅ SOC {len(socs)}건 생성 완료 (총 소요 {_t.time() - t0:.1f}초)"})
            for s in socs:
                tag = " [FOREIGN]" if s.get("is_foreign") else ""
                yield _sse("step", {"phase": "done", "message": f"  · {s['reinsurer']} — {s['share']*100:.2f}% / {s['amount_fmt']} [{s.get('confidence','?')}]{tag}"})
    
            # Mark evidence rows as "used" if their reinsurer matches a resulting SOC
            try:
                from app.routers.soc_test_suite import _norm_reinsurer
                used_reinsurers = {_norm_reinsurer(s["reinsurer"]) for s in socs if s.get("reinsurer")}
                for d in past_claims_data:
                    d["used"] = bool(d.get("reinsurer")) and _norm_reinsurer(d["reinsurer"]) in used_reinsurers
                for d in contracts_data:
                    d["used"] = bool(d.get("reinsurer")) and _norm_reinsurer(d["reinsurer"]) in used_reinsurers
            except Exception:
                pass
    
            # Final result
            evidence_payload = {
                "file_reinsurer": {
                    "name": ai_result.get("file_reinsurer_name", ""),
                    "share": file_share,
                    "amount": file_amount,
                },
                "past_claims": {"search_method": claims_method, "rows": past_claims_data},
                "contracts": {"search_method": contracts_method, "rows": contracts_data},
                "policy_match": policy_data,
                # Co-insurance fan-out breakdown — populated when the
                # underlying data shows >= 2 (cedant, reinsurer) pairs
                # from past claims and/or active contracts. The frontend
                # uses this to display a "co-insurance fan-out" banner
                # explaining why N slips were generated and from which
                # data source. Field name kept as "multi_cedant" for
                # backward compat with already-saved cases.
                "multi_cedant": {
                    "active": multi_cedant,
                    "cedants": distinct_cedants,
                    "groups": [
                        {
                            "cedant": ced,
                            "count": sum(
                                g["count"]
                                for (c, _), g in fanout_groups.items()
                                if c == ced
                            ),
                            "reinsurers": [
                                {
                                    "name": rn,
                                    "count": fanout_groups[(c, rn)]["count"],
                                    "avg_share": (
                                        fanout_groups[(c, rn)]["sum_share"]
                                        / fanout_groups[(c, rn)]["count"]
                                        if fanout_groups[(c, rn)]["count"]
                                        else 0
                                    ),
                                    "sum_amt": fanout_groups[(c, rn)]["sum_amt"],
                                    "source": fanout_groups[(c, rn)]["source"],
                                }
                                for (c, rn) in fanout_groups
                                if c == ced
                            ],
                        }
                        for ced in distinct_cedants
                    ],
                } if multi_cedant else None,
            }
            # Phase 3c: invoke Phase-3d hook (no-op transitions=[]) and
            # attach kb_audit to evidence_payload. OFF mode is a strict
            # no-op so legacy clients see byte-identical evidence.
            finalize_kb_integration(kb_ctx, evidence_payload)
            yield _sse("result", {
                "id": str(case_id),
                "extracted": extracted,
                "policy": policy_data,
                "socs": socs,
                "thinking": thinking,
                "evidence": evidence_payload,
            })

            # Persist final result to the case row so the permalink works.
            try:
                case.status = "done"
                case.extracted = extracted
                case.evidence = evidence_payload
                case.thinking = thinking
                case.policy_match = policy_data
                case.socs = socs
                await db.merge(case)
                await db.commit()
            except Exception as e:
                logger.exception("slip_case persist failed: %s", e)

            yield _sse("done", {})

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache, no-transform",
            "X-Accel-Buffering": "no",
            "Connection": "keep-alive",
        },
    )


# ─── Case permalink endpoints ──────────────────────────────────────────────

from datetime import datetime as _dt
from fastapi import HTTPException, Depends
from pydantic import BaseModel
from sqlalchemy import select as _select
from app.database import get_db as _get_db


class SlipCasePatch(BaseModel):
    doc_type: str | None = None
    edited_socs: list | None = None
    verdict: str | None = None
    # Bordereau-table column layout — see ``SlipCase.columns_config``.
    # ``None`` on the patch ⇒ "leave existing value alone"; pass
    # ``[]`` explicitly to reset to defaults on the frontend.
    columns_config: list | None = None


def _serialize_case(c: SlipCase) -> dict:
    return {
        "id": str(c.id),
        "doc_type": c.doc_type,
        "status": c.status,
        "extracted": c.extracted,
        "evidence": c.evidence,
        "thinking": c.thinking,
        "policy_match": c.policy_match,
        "socs": c.socs,
        "edited_socs": c.edited_socs,
        "columns_config": c.columns_config,
        "error_message": c.error_message,
        "last_edited_at": c.last_edited_at.isoformat() if c.last_edited_at else None,
        "edit_count": c.edit_count,
        "reviewed_at": c.reviewed_at.isoformat() if c.reviewed_at else None,
        "verdict": c.verdict,
        "created_at": c.created_at.isoformat() if c.created_at else None,
        "updated_at": c.updated_at.isoformat() if c.updated_at else None,
    }


@router.get("/case/{case_id}")
async def get_slip_case(case_id: str, db: AsyncSession = Depends(_get_db)):
    try:
        uid = _uuid.UUID(case_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="invalid uuid")
    c = (await db.execute(_select(SlipCase).where(SlipCase.id == uid))).scalar_one_or_none()
    if not c or c.is_deleted:
        raise HTTPException(status_code=404, detail="not found")
    return _serialize_case(c)


@router.patch("/case/{case_id}")
async def patch_slip_case(
    case_id: str,
    patch: SlipCasePatch,
    db: AsyncSession = Depends(_get_db),
):
    try:
        uid = _uuid.UUID(case_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="invalid uuid")
    c = (await db.execute(_select(SlipCase).where(SlipCase.id == uid))).scalar_one_or_none()
    if not c or c.is_deleted:
        raise HTTPException(status_code=404, detail="not found")

    if patch.doc_type is not None:
        c.doc_type = patch.doc_type
    if patch.edited_socs is not None:
        c.edited_socs = patch.edited_socs
        c.last_edited_at = _dt.utcnow()
        c.edit_count = (c.edit_count or 0) + 1
    if patch.columns_config is not None:
        # Empty list is a legitimate value (reset to defaults on FE);
        # only ``None`` means "do not touch", handled by the
        # outer-condition above.
        c.columns_config = patch.columns_config
    if patch.verdict is not None:
        c.verdict = patch.verdict
        c.reviewed_at = _dt.utcnow()

    await db.commit()
    await db.refresh(c)
    return _serialize_case(c)


@router.get("/cases")
async def list_slip_cases(limit: int = 50, db: AsyncSession = Depends(_get_db)):
    rows = (
        await db.execute(
            _select(SlipCase)
            .where(SlipCase.is_deleted.is_(False))
            .order_by(SlipCase.created_at.desc())
            .limit(max(1, min(limit, 200)))
        )
    ).scalars().all()
    return [_serialize_case(c) for c in rows]


class TranslatePayload(BaseModel):
    """Korean → English on-demand translator input. All fields optional —
    the prompt tolerates empty strings and the response preserves the
    same keys so the frontend can merge with whatever subset it sends."""
    account_name: str | None = None
    line: str | None = None
    location_of_loss: str | None = None
    nature_of_loss: str | None = None
    particulars: str | None = None
    remarks: str | None = None
    description: str | None = None


@router.post("/translate")
async def translate_slip_text(payload: TranslatePayload):
    """Translate Korean slip text fields to English. Powers the KR/EN
    toggle on the Claim Detail card for domestic-only cases, where the
    main stream pipeline skips translation because no foreign reinsurer
    is in the SOC mix. Returns the same shape as extracted.text_en."""
    genai.configure(api_key=settings.gemini_api_key)
    model = genai.GenerativeModel("gemini-3.1-flash-lite")
    prompt = SLIP_TRANSLATE_TO_ENGLISH.format(
        account_name=payload.account_name or "",
        line=payload.line or "",
        location_of_loss=payload.location_of_loss or "",
        nature_of_loss=payload.nature_of_loss or "",
        particulars=payload.particulars or payload.description or "",
        remarks=payload.remarks or "",
        description=payload.description or "",
    )
    try:
        resp = await asyncio.to_thread(model.generate_content, prompt)
        text_en = _parse_json(resp.text)
        if not isinstance(text_en, dict):
            raise HTTPException(
                status_code=502,
                detail="translation model returned non-JSON response",
            )
        return text_en
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("on-demand translation failed")
        raise HTTPException(status_code=502, detail=f"translation failed: {e}")
