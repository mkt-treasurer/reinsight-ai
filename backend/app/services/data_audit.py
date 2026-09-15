"""Audit: compare Actual data files with database records."""

import logging
from pathlib import Path

import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.contract import Contract

logger = logging.getLogger(__name__)
DATA_DIR = Path("/app/data")
ACTUAL_DIR = DATA_DIR / "보험료" / "Actual2026 3월 자료"


def _parse_stmt(file_path: Path) -> dict | None:
    """Parse a STMT xlsx and extract cover_note_no + reinsurer amounts."""
    try:
        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(max_row=30, values_only=True))
        wb.close()

        result = {"file": file_path.name, "cover_note_no": None, "assured": None, "cedant": None, "reinsurers": [], "total_premium": None}

        for row in rows:
            vals = [v for v in row if v is not None]
            s = " ".join(str(v) for v in vals)

            # Extract cover note
            if "Cover Note" in s:
                for i, v in enumerate(vals):
                    if str(v).strip() == "Cover Note No.":
                        if i + 1 < len(vals):
                            result["cover_note_no"] = str(vals[i + 1]).strip()
                    if str(v).strip() == "Assured":
                        if i + 1 < len(vals):
                            result["assured"] = str(vals[i + 1]).strip()
                    if str(v).strip() in ("Org. Insurer", "Cedant"):
                        if i + 1 < len(vals):
                            result["cedant"] = str(vals[i + 1]).strip()

            # Extract reinsurer rows (numeric index, name, share, amounts)
            if vals and isinstance(vals[0], (int, float)) and vals[0] >= 1 and vals[0] <= 50:
                if len(vals) >= 5:
                    try:
                        reinsurer = str(vals[1]).strip() if len(vals) > 1 else None
                        share = float(vals[2]) if len(vals) > 2 and isinstance(vals[2], (int, float)) else None
                        gross = float(vals[3]) if len(vals) > 3 and isinstance(vals[3], (int, float)) else None
                        commission = float(vals[4]) if len(vals) > 4 and isinstance(vals[4], (int, float)) else None
                        net = float(vals[5]) if len(vals) > 5 and isinstance(vals[5], (int, float)) else None

                        if reinsurer and gross:
                            result["reinsurers"].append({
                                "name": reinsurer,
                                "share": share,
                                "gross_premium": gross,
                                "commission": commission,
                                "net_premium": net,
                            })
                    except (ValueError, TypeError):
                        pass

            # Total row
            if vals and str(vals[0]).strip() == "Total" and len(vals) >= 3:
                for v in vals[1:]:
                    if isinstance(v, (int, float)) and v > 100:
                        result["total_premium"] = float(v)
                        break

        if result["cover_note_no"] and result["reinsurers"]:
            return result
        return None
    except Exception as e:
        logger.warning(f"STMT parse error {file_path.name}: {e}")
        return None


async def audit_actual_vs_db(db: AsyncSession) -> dict:
    """Compare STMT files with database contracts."""
    import glob

    stmt_files = [Path(f) for f in glob.glob(str(ACTUAL_DIR / "**" / "STMT*.xlsx"), recursive=True)]
    logger.info(f"Found {len(stmt_files)} STMT files to audit")

    matches = []
    mismatches = []
    not_found = []
    parse_errors = []

    for f in stmt_files:
        parsed = _parse_stmt(f)
        if not parsed:
            parse_errors.append({"file": f.name, "reason": "parse_failed"})
            continue

        cn = parsed["cover_note_no"]

        # Find matching contracts in DB
        db_contracts = (await db.execute(
            select(Contract).where(Contract.cover_note_no == cn)
        )).scalars().all()

        if not db_contracts:
            not_found.append({
                "file": f.name,
                "cover_note_no": cn,
                "assured": parsed["assured"],
                "stmt_reinsurers": len(parsed["reinsurers"]),
                "stmt_total": parsed["total_premium"],
            })
            continue

        # Compare each reinsurer
        db_by_reinsurer = {}
        for c in db_contracts:
            key = (c.reinsurer or "").strip()
            db_by_reinsurer[key] = {
                "ri_prem": float(c.ri_prem) if c.ri_prem else 0,
                "ri_commission": float(c.ri_commission) if c.ri_commission else 0,
                "net_ri_prem": float(c.net_ri_prem) if c.net_ri_prem else 0,
                "share": float(c.share) if c.share else 0,
            }

        file_result = {
            "file": f.name,
            "cover_note_no": cn,
            "assured": parsed["assured"],
            "reinsurer_diffs": [],
            "all_match": True,
        }

        for stmt_ri in parsed["reinsurers"]:
            ri_name = stmt_ri["name"]
            # Try to find matching DB reinsurer (fuzzy)
            db_match = None
            for db_name, db_vals in db_by_reinsurer.items():
                if db_name.lower() in ri_name.lower() or ri_name.lower() in db_name.lower():
                    db_match = (db_name, db_vals)
                    break

            if not db_match:
                file_result["reinsurer_diffs"].append({
                    "reinsurer": ri_name,
                    "status": "not_in_db",
                    "stmt_gross": stmt_ri["gross_premium"],
                })
                file_result["all_match"] = False
                continue

            db_name, db_vals = db_match
            # Compare amounts
            diffs = {}
            stmt_gross = stmt_ri["gross_premium"] or 0
            db_gross = db_vals["ri_prem"]
            if abs(stmt_gross - db_gross) > 1:
                diffs["gross_premium"] = {"stmt": stmt_gross, "db": db_gross, "diff": stmt_gross - db_gross}

            stmt_comm = stmt_ri["commission"] or 0
            db_comm = db_vals["ri_commission"]
            if abs(stmt_comm - db_comm) > 1:
                diffs["commission"] = {"stmt": stmt_comm, "db": db_comm, "diff": stmt_comm - db_comm}

            stmt_net = stmt_ri["net_premium"] or 0
            db_net = db_vals["net_ri_prem"]
            if abs(stmt_net - db_net) > 1:
                diffs["net_premium"] = {"stmt": stmt_net, "db": db_net, "diff": stmt_net - db_net}

            if diffs:
                file_result["reinsurer_diffs"].append({
                    "reinsurer": ri_name,
                    "db_reinsurer": db_name,
                    "status": "mismatch",
                    "diffs": diffs,
                })
                file_result["all_match"] = False
            else:
                file_result["reinsurer_diffs"].append({
                    "reinsurer": ri_name,
                    "db_reinsurer": db_name,
                    "status": "match",
                })

        if file_result["all_match"]:
            matches.append(file_result)
        else:
            mismatches.append(file_result)

    return {
        "total_files": len(stmt_files),
        "parsed": len(stmt_files) - len(parse_errors),
        "matches": len(matches),
        "mismatches": len(mismatches),
        "not_found_in_db": len(not_found),
        "parse_errors": len(parse_errors),
        "mismatch_details": mismatches,
        "not_found_details": not_found,
        "match_details": matches,
    }
