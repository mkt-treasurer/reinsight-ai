"""Parse documents and extract structured data using file parsers + Gemini."""

import json
import logging
import re
from pathlib import Path

import google.generativeai as genai
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.models.policy import PolicyDocument

logger = logging.getLogger(__name__)
DATA_DIR = Path("/app/data")

EXTRACT_PROMPT = """You are a reinsurance document analyst. Extract key data from this document content.

Return JSON with these fields (use null if not found):
{{
  "assured": "피보험자/insured name",
  "reinsurer": "재보험사 name(s)",
  "cedant": "출재사",
  "line": "보험종목 (Property/Casualty/Marine/Engineering etc)",
  "currency": "통화",
  "premium": "보험료 금액 (number)",
  "commission_rate": "수수료율 (decimal, e.g. 0.20)",
  "commission": "수수료 금액 (number)",
  "share": "지분율 (decimal, e.g. 0.10 for 10%)",
  "period_from": "보험기간 시작 (YYYY-MM-DD)",
  "period_to": "보험기간 종료 (YYYY-MM-DD)",
  "debit_amount": "차변 금액",
  "credit_amount": "대변 금액",
  "net_amount": "순금액",
  "loss_date": "사고일 (YYYY-MM-DD)",
  "claim_amount": "보험금 금액",
  "ref_no": "참조번호",
  "summary": "문서 요약 (Korean, 1-2 sentences)",
  "email_from": "발신자 (for emails)",
  "email_to": "수신자 (for emails)",
  "email_date": "발신일 (for emails)",
  "email_subject": "제목 (for emails)"
}}

Only return JSON. Set fields to null if not applicable to this document type."""


def _parse_msg(file_path: Path) -> str | None:
    """Extract text from Outlook .msg file."""
    try:
        import extract_msg
        msg = extract_msg.Message(str(file_path))
        parts = []
        if msg.subject:
            parts.append(f"Subject: {msg.subject}")
        if msg.sender:
            parts.append(f"From: {msg.sender}")
        if msg.to:
            parts.append(f"To: {msg.to}")
        if msg.date:
            parts.append(f"Date: {msg.date}")
        if msg.body:
            parts.append(f"\n{msg.body[:3000]}")
        msg.close()
        return "\n".join(parts) if parts else None
    except Exception as e:
        logger.warning(f"MSG parse error {file_path.name}: {e}")
        return None


def _parse_pdf(file_path: Path) -> str | None:
    """Extract text from PDF."""
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(str(file_path))
        texts = []
        for page in reader.pages[:10]:  # Max 10 pages
            t = page.extract_text()
            if t:
                texts.append(t)
        return "\n".join(texts)[:5000] if texts else None
    except Exception as e:
        logger.warning(f"PDF parse error {file_path.name}: {e}")
        return None


def _parse_docx(file_path: Path) -> str | None:
    """Extract text from .docx."""
    try:
        from docx import Document
        doc = Document(str(file_path))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also get tables
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    texts.append(" | ".join(cells))
        return "\n".join(texts)[:5000] if texts else None
    except Exception as e:
        logger.warning(f"DOCX parse error {file_path.name}: {e}")
        return None


def _parse_xlsx(file_path: Path) -> str | None:
    """Extract summary from Excel."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        texts = []
        for sheet_name in wb.sheetnames[:3]:
            ws = wb[sheet_name]
            texts.append(f"[Sheet: {sheet_name}]")
            for i, row in enumerate(ws.iter_rows(max_row=20, values_only=True)):
                vals = [str(v) for v in row if v is not None]
                if vals:
                    texts.append(" | ".join(vals))
        wb.close()
        return "\n".join(texts)[:5000] if texts else None
    except Exception as e:
        logger.warning(f"XLSX parse error {file_path.name}: {e}")
        return None


def _parse_file(file_path: Path) -> str | None:
    """Parse any supported file type and return text content."""
    suffix = file_path.suffix.lower()
    if suffix == ".msg":
        return _parse_msg(file_path)
    elif suffix == ".pdf":
        return _parse_pdf(file_path)
    elif suffix in (".docx",):
        return _parse_docx(file_path)
    elif suffix in (".xlsx",):
        return _parse_xlsx(file_path)
    elif suffix in (".doc", ".xls"):
        # Legacy formats - can't easily parse without extra deps
        return None
    return None


def _extract_with_gemini(text: str, file_name: str, doc_type: str) -> dict | None:
    """Use Gemini to extract structured data from document text."""
    try:
        genai.configure(api_key=settings.gemini_api_key)
        model = genai.GenerativeModel("gemini-3.1-flash-lite")

        prompt = f"{EXTRACT_PROMPT}\n\nDocument type: {doc_type}\nFile name: {file_name}\n\nContent:\n{text[:4000]}"
        response = model.generate_content(prompt)
        raw = response.text.strip()
        if "```" in raw:
            raw = raw.split("```")[1]
            if raw.startswith("json"):
                raw = raw[4:]
            raw = raw.strip()
        return json.loads(raw)
    except Exception as e:
        logger.warning(f"Gemini extract error {file_name}: {e}")
        return None


async def parse_documents(db: AsyncSession, limit: int = 50, doc_type: str | None = None) -> dict:
    """Parse unparsed documents and extract data with Gemini."""
    query = select(PolicyDocument).where(PolicyDocument.extracted_data.is_(None))
    if doc_type:
        query = query.where(PolicyDocument.doc_type == doc_type)
    query = query.limit(limit)

    docs = (await db.execute(query)).scalars().all()

    parsed = 0
    failed = 0
    skipped = 0

    for doc in docs:
        file_path = DATA_DIR / doc.file_path
        if not file_path.exists():
            skipped += 1
            continue

        # Parse file content
        text = _parse_file(file_path)
        if not text:
            # Mark as unparsable so we don't retry
            doc.extracted_data = {"_status": "unparsable", "_reason": f"unsupported format: {file_path.suffix}"}
            failed += 1
            continue

        # Extract with Gemini
        extracted = _extract_with_gemini(text, doc.file_name, doc.doc_type or "unknown")
        if extracted:
            extracted["_status"] = "parsed"
            extracted["_text_length"] = len(text)
            doc.extracted_data = extracted
            parsed += 1
        else:
            doc.extracted_data = {"_status": "extract_failed", "_text_length": len(text)}
            failed += 1

    await db.commit()
    logger.info(f"Parsed {parsed} documents, {failed} failed, {skipped} skipped")
    return {"parsed": parsed, "failed": failed, "skipped": skipped, "total": len(docs)}
