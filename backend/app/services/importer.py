import logging
from datetime import datetime
from pathlib import Path

import openpyxl
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.contract import Contract
from app.models.claim import Claim
from app.models.cover_note import CoverNote
from app.models.company import Company

logger = logging.getLogger(__name__)

DATA_DIR = Path("/app/data")


def _safe_date(val):
    if isinstance(val, datetime):
        return val.date()
    return None


def _safe_numeric(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return val
    try:
        return float(str(val).replace(",", ""))
    except (ValueError, TypeError):
        return None


def _safe_str(val, max_len: int | None = None):
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    if max_len and len(s) > max_len:
        s = s[:max_len]
    return s


def _cell_map_from_row(row):
    """Build a dict mapping column letters (A, B, ..., AA, AB, ...) to cell values.

    Accepts either a tuple (values_only=True) or a sequence of Cell objects.
    """
    cell_map = {}
    for idx, cell in enumerate(row):
        col_idx = idx + 1
        if col_idx <= 26:
            letter = chr(64 + col_idx)
        else:
            letter = chr(64 + (col_idx - 1) // 26) + chr(65 + (col_idx - 1) % 26)
        val = cell.value if hasattr(cell, "value") else cell
        cell_map[letter] = val
    return cell_map


# Header keyword → canonical field name. Case-insensitive substring match.
# First match wins, so order more-specific keywords first.
_HEADER_MAP = [
    ("cont. month", "cont_month"),
    ("cont.month", "cont_month"),
    ("no.", "no"),
    ("covernote", "cover_note_no"),
    ("cover note", "cover_note_no"),
    ("assured", "assured"),
    ("project name", "project_name"),
    ("line 2", "line2"),
    ("line2", "line2"),
    ("line 1", "line"),
    ("line", "line"),
    ("new/renew", "new_renew"),
    ("renewable", "renewable"),
    ("retail", "retail"),
    ("original cedant", "original_cedant"),
    ("cedant", "cedant"),
    ("from", "period_from"),
    ("to", "period_to"),
    ("installment", "installment"),
    ("ppw", "ppw"),
    ("currency", "currency"),
    ("gross prem.(100", "gross_prem_100"),
    ("gross prem.(ins", "gross_prem_inst"),
    ("reinsurer", "reinsurer"),
    ("share", "share"),
    ("r/i prem", "ri_prem"),
    ("reins. prem", "ri_prem"),
    ("ri/c", "ri_commission"),
    ("net ri/p (kwon)", "net_ri_prem_kwon"),
    ("net ri/p(kwon", "net_ri_prem_kwon"),
    ("net ri/p", "net_ri_prem"),
    ("net to uwr", "net_to_uwr"),
    ("rec. date", "rec_date"),
    ("rec.date", "rec_date"),
    ("paid date", "paid_date"),
    ("co-brokerage", "co_brokerage"),
    ("partner", "partner"),
]


def _build_header_map(ws) -> tuple[dict[str, int], int] | None:
    """Scan first 10 rows for header. Return (field→col_idx map, first_data_row)."""
    for header_row_idx in range(1, 11):
        row = list(next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, max_col=40, values_only=True), []))
        if not row:
            continue
        # Header rows must contain "Assured" and "Reinsurer" somewhere
        texts = [str(v).strip().lower() if v else "" for v in row]
        joined = " | ".join(texts)
        if "assured" in joined and "reinsurer" in joined and ("cover" in joined or "note" in joined):
            # Found header row — map each column
            mapping: dict[str, int] = {}
            for idx, text in enumerate(texts):
                if not text:
                    continue
                for keyword, field in _HEADER_MAP:
                    if keyword in text and field not in mapping:
                        mapping[field] = idx
                        break
            return mapping, header_row_idx + 1
    return None


async def import_contracts(db: AsyncSession, company_id: int = 1):
    """Import AI계약관리.xlsx into contracts table.

    Each Actual* sheet has its own column layout — header scanning is dynamic.
    """
    file_path = DATA_DIR / "보험료" / "AI계약관리.xlsx"
    if not file_path.exists():
        logger.warning(f"Contract file not found: {file_path}")
        return 0

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    count = 0

    for sheet_name in wb.sheetnames:
        if not sheet_name.startswith("Actual"):
            continue

        year_str = sheet_name.replace("Actual", "")
        try:
            year = int(year_str)
        except ValueError:
            continue

        sheet_count = 0
        ws = wb[sheet_name]
        hdr = _build_header_map(ws)
        if not hdr:
            logger.warning(f"  {sheet_name}: header not found, skipping")
            continue
        col_idx, first_data_row = hdr
        max_col = max(col_idx.values()) + 2 if col_idx else 30
        logger.info(f"Importing {sheet_name} (header row {first_data_row - 1}, fields {len(col_idx)})...")

        def get(row, field):
            i = col_idx.get(field)
            return row[i] if i is not None and i < len(row) else None

        for row in ws.iter_rows(min_row=first_data_row, max_col=max_col, values_only=True):
            if not get(row, "cover_note_no") and not get(row, "assured"):
                continue

            contract = Contract(
                company_id=company_id,
                year=year,
                cont_month=_safe_str(get(row, "cont_month"), 50),
                no=_safe_str(get(row, "no"), 50),
                cover_note_no=_safe_str(get(row, "cover_note_no"), 100),
                assured=_safe_str(get(row, "assured"), 500),
                project_name=_safe_str(get(row, "project_name"), 500),
                line=_safe_str(get(row, "line"), 100),
                line2=_safe_str(get(row, "line2"), 100),
                new_renew=_safe_str(get(row, "new_renew"), 20),
                renewable=_safe_str(get(row, "renewable"), 20),
                retail=_safe_str(get(row, "retail"), 20),
                original_cedant=_safe_str(get(row, "original_cedant"), 100),
                cedant=_safe_str(get(row, "cedant"), 100),
                period_from=_safe_date(get(row, "period_from")),
                period_to=_safe_date(get(row, "period_to")),
                installment=_safe_str(get(row, "installment"), 50),
                ppw=_safe_str(get(row, "ppw"), 100),
                currency=_safe_str(get(row, "currency"), 20),
                gross_prem_100=_safe_numeric(get(row, "gross_prem_100")),
                gross_prem_inst=_safe_numeric(get(row, "gross_prem_inst")),
                reinsurer=_safe_str(get(row, "reinsurer"), 100),
                share=_safe_numeric(get(row, "share")),
                ri_prem=_safe_numeric(get(row, "ri_prem")),
                ri_commission=_safe_numeric(get(row, "ri_commission")),
                net_ri_prem=_safe_numeric(get(row, "net_ri_prem")),
                net_ri_prem_kwon=_safe_numeric(get(row, "net_ri_prem_kwon")),
                net_to_uwr=_safe_numeric(get(row, "net_to_uwr")),
                rec_date=_safe_date(get(row, "rec_date")),
                paid_date=_safe_date(get(row, "paid_date")),
                co_brokerage=_safe_numeric(get(row, "co_brokerage")),
                partner=_safe_str(get(row, "partner"), 200),
            )
            rec = _safe_date(get(row, "rec_date"))
            paid = _safe_date(get(row, "paid_date"))
            if paid:
                contract.workflow_status = "completed"
            elif rec:
                contract.workflow_status = "payment_received"
            elif contract.reinsurer:
                contract.workflow_status = "sent_to_reinsurer"
            else:
                contract.workflow_status = "booked"
            db.add(contract)
            count += 1
            sheet_count += 1

        # Commit per sheet — avoids one giant transaction + makes progress visible
        try:
            await db.commit()
            logger.info(f"  {sheet_name}: +{sheet_count} (total {count})")
        except Exception as e:
            await db.rollback()
            logger.error(f"  {sheet_name} commit failed: {e}")

    wb.close()
    logger.info(f"Imported {count} contracts")
    return count


async def import_claims(db: AsyncSession, company_id: int = 1):
    """Import Claim list(통합) & pending List.xlsx into claims table."""
    file_path = DATA_DIR / "보험금" / "Claim list(통합) & pending List.xlsx"
    if not file_path.exists():
        logger.warning(f"Claims file not found: {file_path}")
        return 0

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    count = 0

    if "ALL" not in wb.sheetnames:
        wb.close()
        return 0

    ws = wb["ALL"]
    for row in ws.iter_rows(min_row=3, values_only=False):
        cell_map = _cell_map_from_row(row)

        if not cell_map.get("E") and not cell_map.get("I"):
            continue

        claim = Claim(
            company_id=company_id,
            sheet_name="ALL",
            booking_month=_safe_str(cell_map.get("B")),
            soc_received=_safe_str(cell_map.get("C")),
            soc_sent=_safe_str(cell_map.get("D")),
            account_name=_safe_str(cell_map.get("E")),
            line=_safe_str(cell_map.get("F")),
            policy_period=_safe_date(cell_map.get("G")),
            dol=_safe_date(cell_map.get("H")),
            reinsurer=_safe_str(cell_map.get("I")),
            currency=_safe_str(cell_map.get("J")),
            total_amount=_safe_numeric(cell_map.get("K")),
            share=_safe_numeric(cell_map.get("L")),
            origin_currency=_safe_numeric(cell_map.get("M")),
            roe=_safe_numeric(cell_map.get("N")),
            krw_amount=_safe_numeric(cell_map.get("O")),
            soc_amount=_safe_numeric(cell_map.get("P")),
            cedant=_safe_str(cell_map.get("Q")),
            status=_safe_str(cell_map.get("R")),
            received_date=_safe_date(cell_map.get("S")),
            paid_date=_safe_date(cell_map.get("T")),
            account_mgr=_safe_str(cell_map.get("U")),
            remarks=_safe_str(cell_map.get("V")),
            ref_no=_safe_str(cell_map.get("W")),
        )
        # Auto-infer workflow status
        cl_status = _safe_str(cell_map.get("R"))
        cl_paid = _safe_date(cell_map.get("T"))
        cl_received = _safe_date(cell_map.get("S"))
        cl_sent = _safe_str(cell_map.get("D"))
        if cl_status == "Closed" or cl_paid:
            claim.workflow_status = "completed"
        elif cl_received:
            claim.workflow_status = "payment_received"
        elif cl_sent and cl_sent.strip() not in ("", "미 발송"):
            claim.workflow_status = "sent_to_reinsurer"
        elif _safe_str(cell_map.get("C")):
            claim.workflow_status = "soc_received"
        else:
            claim.workflow_status = "booked"
        db.add(claim)
        count += 1

        if count % 500 == 0:
            await db.flush()

    await db.commit()
    wb.close()
    logger.info(f"Imported {count} claims")
    return count


async def import_cover_notes(db: AsyncSession, company_id: int = 1):
    """Import Cover Note.xlsx into cover_notes table."""
    file_path = DATA_DIR / "보험료" / "Cover Note.xlsx"
    if not file_path.exists():
        logger.warning(f"Cover note file not found: {file_path}")
        return 0

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    count = 0

    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(min_row=4, values_only=False):
        cell_map = _cell_map_from_row(row)

        if not cell_map.get("B"):
            continue

        cn = CoverNote(
            company_id=company_id,
            issuing_date=_safe_date(cell_map.get("A")),
            cover_note_number=_safe_str(cell_map.get("B")),
            assured=_safe_str(cell_map.get("C")),
            reassured=_safe_str(cell_map.get("D")),
            line=_safe_str(cell_map.get("E")),
            account=_safe_str(cell_map.get("F")),
            remarks=_safe_str(cell_map.get("G")),
        )
        db.add(cn)
        count += 1

    await db.commit()
    wb.close()
    logger.info(f"Imported {count} cover notes")
    return count


async def import_all(db: AsyncSession):
    """Run full import. Clears existing data first."""
    # Check if company exists, create if not
    result = await db.execute(select(Company).where(Company.code == "ins-insurance"))
    company = result.scalar_one_or_none()
    if not company:
        company = Company(name="인스보험", code="ins-insurance")
        db.add(company)
        await db.commit()
        await db.refresh(company)

    company_id = company.id

    await db.execute(text("TRUNCATE policy_claim_links, policies, contract_claim_links, contracts, claims, cover_notes RESTART IDENTITY CASCADE"))
    await db.commit()

    contracts = await import_contracts(db, company_id)
    claims = await import_claims(db, company_id)
    cover_notes = await import_cover_notes(db, company_id)

    return {"contracts": contracts, "claims": claims, "cover_notes": cover_notes}
